"""Tests for the real TaPEx baseline.

We exercise TaPExBaseline against an in-memory pipeline mock so the test is
hermetic — no model download, no network, no GPU. The fact that the wrapper
correctly reads the workbook range, builds the column-major dict TaPEx
expects, and wraps the answer in `[...]` is what we verify.
"""
from __future__ import annotations

import os
import shutil
import tempfile
import unittest

import openpyxl

from baselines import (
    BINDER_UNAVAILABLE_REASON,
    BaselineUnavailable,
    BinderBaseline,
    TaPExBaseline,
    _read_table,
)


class FakePipeline:
    """Stand-in for ``transformers.pipeline('table-question-answering', ...)``."""
    def __init__(self, response_answer: str = "42"):
        self.calls = []
        self.response_answer = response_answer

    def __call__(self, table=None, query=None, **kwargs):
        self.calls.append({"table": table, "query": query, "kwargs": kwargs})
        return {"answer": self.response_answer}


class TestReadTable(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="sllm_tapex_")
        self.path = os.path.join(self.tmp, "wb.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        ws['A1'] = "ID"
        ws['B1'] = "Name"
        ws['C1'] = "Score"
        ws['A2'] = 1
        ws['B2'] = "Alice"
        ws['C2'] = 90
        ws['A3'] = 2
        ws['B3'] = "Bob"
        ws['C3'] = 85
        wb.save(self.path)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_reads_header_and_body(self):
        header, body = _read_table(self.path, "Sheet", "A1:C3")
        self.assertEqual(header, ["ID", "Name", "Score"])
        self.assertEqual(body, [["1", "Alice", "90"], ["2", "Bob", "85"]])

    def test_synthesizes_header_when_first_row_is_numeric(self):
        # No string header — TaPEx would choke; the wrapper synthesizes A,B,C.
        wb = openpyxl.load_workbook(self.path)
        ws = wb.active
        ws['A1'] = 1
        ws['B1'] = 2
        ws['C1'] = 3
        wb.save(self.path)

        header, body = _read_table(self.path, "Sheet", "A1:C3")
        self.assertEqual(header, ["A", "B", "C"])
        self.assertEqual(len(body), 3)


class TestTaPExBaseline(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="sllm_tapex2_")
        self.path = os.path.join(self.tmp, "wb.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        ws['A1'] = "ID"
        ws['B1'] = "Score"
        ws['A2'] = 1
        ws['B2'] = 90
        ws['A3'] = 2
        ws['B3'] = 85
        wb.save(self.path)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_answer_wraps_in_brackets(self):
        fake = FakePipeline(response_answer="90")
        tapex = TaPExBaseline(pipeline=fake)
        result = tapex.answer(self.path, "Sheet", "A1:B3", "What is Alice's score?")
        self.assertEqual(result, "[90]")

    def test_pipeline_receives_column_dict(self):
        fake = FakePipeline(response_answer="x")
        tapex = TaPExBaseline(pipeline=fake)
        tapex.answer(self.path, "Sheet", "A1:B3", "Q?")

        self.assertEqual(len(fake.calls), 1)
        call = fake.calls[0]
        self.assertEqual(call["query"], "Q?")
        # TaPEx expects {column: [cells]} — verify the shape.
        table = call["table"]
        self.assertIsInstance(table, dict)
        self.assertIn("ID", table)
        self.assertIn("Score", table)
        self.assertEqual(table["ID"], ["1", "2"])
        self.assertEqual(table["Score"], ["90", "85"])

    def test_handles_list_response(self):
        # HF pipeline sometimes returns a list of dicts.
        class ListPipeline:
            def __call__(self, table=None, query=None, **kw):
                return [{"answer": "first"}, {"answer": "second"}]

        tapex = TaPExBaseline(pipeline=ListPipeline())
        result = tapex.answer(self.path, "Sheet", "A1:B3", "Q?")
        self.assertEqual(result, "[first]")

    def test_disambiguates_duplicate_headers(self):
        # Two columns named "Score" — TaPEx requires unique keys.
        wb = openpyxl.load_workbook(self.path)
        ws = wb.active
        ws['B1'] = "Score"
        # Add a third column also named "Score"
        ws['C1'] = "Score"
        ws['C2'] = 100
        ws['C3'] = 95
        wb.save(self.path)

        fake = FakePipeline(response_answer="95")
        tapex = TaPExBaseline(pipeline=fake)
        tapex.answer(self.path, "Sheet", "A1:C3", "Q?")
        table = fake.calls[0]["table"]
        # Should have 3 distinct keys: Score and Score_1 (and ID).
        self.assertEqual(len(table), 3)
        self.assertIn("Score", table)
        self.assertIn("Score_1", table)

    def test_missing_transformers_raises_clear_error(self):
        """When no pipeline is injected and transformers isn't importable,
        the wrapper raises a RuntimeError with a clear pip hint."""
        import sys
        import unittest.mock as mock
        # Force the lazy import to fail.
        with mock.patch.dict(sys.modules, {"transformers": None}):
            tapex = TaPExBaseline()
            with self.assertRaises(RuntimeError) as cm:
                tapex.answer(self.path, "Sheet", "A1:B3", "Q?")
            self.assertIn("transformers", str(cm.exception).lower())


class TestBinderBaseline(unittest.TestCase):

    def test_binder_is_explicitly_unavailable(self):
        binder = BinderBaseline()

        self.assertEqual(binder.status, "unavailable")
        self.assertEqual(
            binder.skip_reason(),
            {"component": "binder", "reason": BINDER_UNAVAILABLE_REASON},
        )
        with self.assertRaises(BaselineUnavailable) as cm:
            binder.answer("workbook.xlsx", "Sheet", "A1:B2", "Q?")
        self.assertEqual(cm.exception.baseline, "Binder")
        self.assertEqual(cm.exception.reason, BINDER_UNAVAILABLE_REASON)


if __name__ == "__main__":
    unittest.main()
