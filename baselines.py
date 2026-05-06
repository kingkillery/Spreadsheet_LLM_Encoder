"""Real baselines for the SpreadsheetLLM QA evaluation.

The paper compares against TaPEx (Liu et al., 2022) and Binder (Cheng et al.,
2023). This module provides a real TaPEx wrapper using HuggingFace
``transformers``. Binder remains TODO — its neural-symbolic SQL execution
loop needs a separate vendoring effort.

The wrapper is intentionally narrow: it loads the model lazily, extracts the
table from a workbook + table-range, runs the pipeline, and wraps the answer
in ``[...]`` to match the paper's QA contract. Network/disk access happens
only when the wrapper is actually called — instantiation is cheap.
"""
from __future__ import annotations

import logging
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

import paper_serializers

logger = logging.getLogger(__name__)


def _read_table(
    workbook_path: str,
    sheet_name: str,
    table_range: str,
) -> Tuple[List[str], List[List[str]]]:
    """Read ``table_range`` of ``sheet_name`` and return ``(header, rows)``.

    The first retained row is treated as the header. Empty cells become
    empty strings; numbers are stringified. Only cell values are returned —
    formulas/styles are dropped because TaPEx operates on textual tables.
    """
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet {sheet_name!r} not in {workbook_path}")
    sheet = wb[sheet_name]
    r1, c1, r2, c2 = paper_serializers.parse_range(table_range)
    paper_serializers._check_range_size(r1, c1, r2, c2, "TaPExBaseline._read_table")

    rows: List[List[str]] = []
    for r in range(r1, r2 + 1):
        row = []
        for c in range(c1, c2 + 1):
            v = sheet.cell(row=r, column=c).value
            row.append("" if v is None else str(v))
        rows.append(row)

    if not rows:
        return [], []

    # Synthesize a header if the first row is all-numeric (TaPEx requires
    # string column names; numeric headers cause its tokenizer to choke).
    header = rows[0]
    body = rows[1:]
    if not header or all(_looks_numeric(c) for c in header if c):
        header = [get_column_letter(c) for c in range(c1, c2 + 1)]
        body = rows
    return header, body


def _looks_numeric(s: str) -> bool:
    if not s:
        return False
    try:
        float(s)
        return True
    except ValueError:
        return False


class TaPExBaseline:
    """HF TaPEx adapter for spreadsheet QA.

    Usage::

        from baselines import TaPExBaseline
        tapex = TaPExBaseline()  # lazy; nothing loaded yet
        answer = tapex.answer(workbook_path, sheet_name, table_range, query)

    The model (``microsoft/tapex-base-finetuned-wtq`` by default) is loaded
    on the first ``.answer(...)`` call. Pass ``pipeline=...`` to inject a
    pre-built or mocked pipeline (used in tests).
    """

    def __init__(
        self,
        model: str = "microsoft/tapex-base-finetuned-wtq",
        pipeline: Any = None,
        max_rows: int = 64,
    ) -> None:
        self.model = model
        self._pipeline = pipeline
        self.max_rows = max_rows

    def _ensure_pipeline(self) -> Any:
        if self._pipeline is not None:
            return self._pipeline
        try:
            from transformers import pipeline as hf_pipeline  # type: ignore
        except ImportError as exc:  # pragma: no cover - environmental
            raise RuntimeError(
                "transformers is not installed; pip install transformers "
                "or pass a pre-built `pipeline=` to TaPExBaseline."
            ) from exc
        self._pipeline = hf_pipeline("table-question-answering", model=self.model)
        return self._pipeline

    def answer(
        self,
        workbook_path: str,
        sheet_name: str,
        table_range: str,
        query: str,
    ) -> str:
        """Run TaPEx on the given range/query. Returns ``[<answer>]``."""
        header, body = _read_table(workbook_path, sheet_name, table_range)
        if not header:
            return "[]"

        # TaPEx wants {column_name: [cells_in_column]}. Keep only the first
        # ``max_rows`` body rows to avoid blowing the model's context.
        body = body[: self.max_rows]

        # Disambiguate duplicate header strings — TaPEx requires unique keys.
        seen_keys: Dict[str, int] = {}
        norm_header: List[str] = []
        for h in header:
            n = seen_keys.get(h, 0)
            seen_keys[h] = n + 1
            norm_header.append(h if n == 0 else f"{h}_{n}")

        table: Dict[str, List[str]] = {h_norm: [] for h_norm in norm_header}
        for row in body:
            for h_norm, v in zip(norm_header, row):
                table[h_norm].append(v)

        pipe = self._ensure_pipeline()
        result = pipe(table=table, query=query)
        # HF returns either a dict or a list of dicts depending on input.
        if isinstance(result, list):
            result = result[0] if result else {}
        ans = result.get("answer", "") if isinstance(result, dict) else str(result)
        return f"[{ans}]"


__all__ = ["TaPExBaseline"]
