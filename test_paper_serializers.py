"""Tests for paper_serializers.py — address helpers, coord remapping, and prompt serializers."""
import json
import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
import paper_serializers as ps


def _fmt_key(sem_type: str, nfs: str = "General") -> str:
    return json.dumps({"type": sem_type, "nfs": nfs}, sort_keys=True)


class TestParseFormatRange(unittest.TestCase):

    def test_check_range_size_rejects_pathological_range(self):
        """Pathological LLM-controlled ranges must raise rather than pin a CPU."""
        with self.assertRaises(ValueError):
            ps._check_range_size(1, 1, 100_000, 100, "test")

    def test_check_range_size_accepts_normal_range(self):
        ps._check_range_size(1, 1, 100, 10, "test")  # ~1k cells, fine.

    def test_parse_range_warns_on_reversed_endpoints(self):
        with self.assertLogs("paper_serializers", level="WARNING") as cm:
            r1, c1, r2, c2 = ps.parse_range("C3:A1")
        self.assertEqual((r1, c1, r2, c2), (1, 1, 3, 3))
        self.assertTrue(any("reversed" in line for line in cm.output))

    def test_single_ref_parses_to_same_row_col_twice(self):
        r1, c1, r2, c2 = ps.parse_range("B3")
        self.assertEqual((r1, c1, r2, c2), (3, 2, 3, 2))

    def test_range_parses_correctly(self):
        r1, c1, r2, c2 = ps.parse_range("A1:C5")
        self.assertEqual((r1, c1, r2, c2), (1, 1, 5, 3))

    def test_parse_format_ref_round_trip(self):
        ref = "D7"
        r, c = ps.split_ref(ref)
        self.assertEqual(ps.format_ref(r, c), ref)

    def test_parse_format_range_round_trip(self):
        rng = "B2:E10"
        r1, c1, r2, c2 = ps.parse_range(rng)
        self.assertEqual(ps.format_range(r1, c1, r2, c2), rng)

    def test_format_range_single_cell_returns_ref_not_range(self):
        self.assertEqual(ps.format_range(3, 2, 3, 2), "B3")

    def test_invalid_ref_raises_value_error(self):
        with self.assertRaises(ValueError):
            ps.split_ref("notaref")


class TestBuildCoordMap(unittest.TestCase):

    def setUp(self):
        self.coord_map = ps.build_coord_map([2, 5, 7], [1, 3])

    def test_row_map_values(self):
        self.assertEqual(self.coord_map["rows"], {2: 1, 5: 2, 7: 3})

    def test_col_map_values(self):
        self.assertEqual(self.coord_map["cols"], {1: 1, 3: 2})

    def test_rows_inv_map(self):
        self.assertEqual(self.coord_map["rows_inv"], {1: 2, 2: 5, 3: 7})

    def test_cols_inv_map(self):
        self.assertEqual(self.coord_map["cols_inv"], {1: 1, 2: 3})

    def test_duplicate_inputs_deduplicated(self):
        cm = ps.build_coord_map([3, 3, 5], [1, 1])
        self.assertEqual(cm["rows"], {3: 1, 5: 2})

    def test_unsorted_inputs_produce_sorted_map(self):
        cm = ps.build_coord_map([7, 2, 5], [3, 1])
        self.assertEqual(cm["rows"], {2: 1, 5: 2, 7: 3})


class TestRemapRange(unittest.TestCase):

    def setUp(self):
        # rows [2,5,7] -> [1,2,3]; cols [1,3] -> [1,2]
        self.coord_map = ps.build_coord_map([2, 5, 7], [1, 3])

    def test_remap_range_single_ref_e5(self):
        # E5 = col E (5) row 5. Col 5 is not in [1,3], so None.
        # Use C5: col C (3) -> compact col 2, row 5 -> compact row 2 => B2
        result = ps.remap_range("C5", self.coord_map)
        self.assertEqual(result, "B2")

    def test_remap_range_returns_none_for_unmapped_row(self):
        # row 4 not in kept_rows
        result = ps.remap_range("A4", self.coord_map)
        self.assertIsNone(result)

    def test_remap_range_returns_none_for_unmapped_col(self):
        # col B (2) not in kept_cols
        result = ps.remap_range("B2", self.coord_map)
        self.assertIsNone(result)

    def test_unremap_range_round_trip(self):
        # C5 -> remap -> B2 -> unremap -> C5
        compact = ps.remap_range("C5", self.coord_map)
        self.assertEqual(compact, "B2")
        original = ps.unremap_range(compact, self.coord_map)
        self.assertEqual(original, "C5")

    def test_remap_range_accepts_json_reloaded_string_keys(self):
        reloaded = json.loads(json.dumps(self.coord_map))
        result = ps.remap_range("C5", reloaded)
        self.assertEqual(result, "B2")

    def test_unremap_range_accepts_json_reloaded_string_keys(self):
        reloaded = json.loads(json.dumps(self.coord_map))
        result = ps.unremap_range("B2", reloaded)
        self.assertEqual(result, "C5")

    def test_unremap_range_returns_none_for_out_of_bounds(self):
        # Compact row 9 does not exist in rows_inv
        result = ps.unremap_range("A9", self.coord_map)
        self.assertIsNone(result)

    def test_remap_range_multi_cell_range(self):
        # A2:A7 = row2->1, row7->3, col1->1 => A1:A3
        result = ps.remap_range("A2:A7", self.coord_map)
        self.assertEqual(result, "A1:A3")


class TestVanillaPrompt(unittest.TestCase):

    def test_2x2_sheet_with_one_empty_cell(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Hello"
        ws["B1"] = "World"
        ws["A2"] = 42
        # B2 is empty
        result = ps.to_paper_vanilla_prompt(ws)
        self.assertEqual(result, "A1,Hello|B1,World|A2,42|B2,")

    def test_pipes_in_values_replaced_with_spaces(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "a|b"
        result = ps.to_paper_vanilla_prompt(ws)
        self.assertIn("A1,a b", result)

    def test_newlines_in_values_replaced_with_spaces(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "line1\nline2"
        result = ps.to_paper_vanilla_prompt(ws)
        self.assertIn("A1,line1 line2", result)

    def test_empty_sheet_emits_only_empty_a1_pair(self):
        # A brand-new openpyxl sheet reports max_row=1, max_column=1 even with
        # no data written, so to_paper_vanilla_prompt emits exactly "A1,".
        wb = openpyxl.Workbook()
        ws = wb.active
        result = ps.to_paper_vanilla_prompt(ws)
        self.assertEqual(result, "A1,")


class TestLabelForFormatKey(unittest.TestCase):

    def test_integer_maps_to_intnum(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("integer", "0")), "IntNum")

    def test_float_maps_to_floatnum(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("float")), "FloatNum")

    def test_date_maps_to_datedata(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("date")), "DateData")

    def test_email_maps_to_emaildata(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("email")), "EmailData")

    def test_percentage_maps_to_percentagenum(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("percentage")), "PercentageNum")

    def test_year_maps_to_paper_year_label(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("year")), "Year")

    def test_informative_integer_nfs_is_emitted(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("integer", "#,##0")), "#,##0")

    def test_informative_date_nfs_is_emitted(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("date", "d-mmm-yy")), "d-mmm-yy")

    def test_informative_time_nfs_is_emitted(self):
        self.assertEqual(ps.label_for_format_key(_fmt_key("time", "H:mm:ss")), "H:mm:ss")

    def test_generic_number_format_is_not_informative(self):
        self.assertFalse(ps.is_informative_number_format("0.00"))

    def test_text_returns_none(self):
        self.assertIsNone(ps.label_for_format_key(_fmt_key("text")))

    def test_boolean_returns_none(self):
        self.assertIsNone(ps.label_for_format_key(_fmt_key("boolean")))

    def test_unknown_type_returns_none(self):
        self.assertIsNone(ps.label_for_format_key(_fmt_key("mystery_type")))

    def test_malformed_json_returns_none(self):
        self.assertIsNone(ps.label_for_format_key("not json {{{"))

    def test_non_string_returns_none(self):
        self.assertIsNone(ps.label_for_format_key(None))  # type: ignore[arg-type]


class TestCompressedPrompt(unittest.TestCase):

    def test_no_formats_emits_literal_value_tuple(self):
        encoding = {"cells": {"42": ["B5"]}, "formats": {}}
        result = ps.to_paper_compressed_prompt(encoding)
        self.assertEqual(result, "(42|B5)")

    def test_compressible_format_region_suppresses_literal_tuples(self):
        # All cells A1:B2 are covered by the integer region
        int_key = _fmt_key("integer", "0")
        encoding = {
            "cells": {"5": ["A1", "A2", "B1", "B2"]},
            "formats": {int_key: ["A1:B2"]},
        }
        result = ps.to_paper_compressed_prompt(encoding)
        self.assertEqual(result, "(IntNum|A1:B2)")

    def test_informative_nfs_region_suppresses_literal_tuples(self):
        date_key = _fmt_key("date", "yyyy/mm/dd")
        encoding = {
            "cells": {"2024-01-01": ["A1", "A2"]},
            "formats": {date_key: ["A1:A2"]},
        }
        result = ps.to_paper_compressed_prompt(encoding)
        self.assertEqual(result, "(yyyy/mm/dd|A1:A2)")

    def test_mixed_literal_and_label_region_in_row_major_order(self):
        int_key = _fmt_key("integer", "0")
        encoding = {
            "cells": {"Header": ["A1"]},
            "formats": {int_key: ["A2:A4"]},
        }
        result = ps.to_paper_compressed_prompt(encoding)
        self.assertIn("(Header|A1)", result)
        self.assertIn("(IntNum|A2:A4)", result)
        # A1 must appear before A2:A4 in row-major order
        self.assertLess(result.index("(Header|A1)"), result.index("(IntNum|A2:A4)"))

    def test_with_coord_map_remaps_addresses(self):
        coord_map = ps.build_coord_map([1, 2], [1, 2])
        int_key = _fmt_key("integer", "0")
        encoding = {
            "cells": {},
            "formats": {int_key: ["A1:B2"]},
        }
        result = ps.to_paper_compressed_prompt(encoding, coord_map=coord_map)
        # A1:B2 remaps to A1:B2 (same compact coords since rows/cols 1,2 map to 1,2)
        self.assertIn("IntNum", result)
        self.assertIn("A1:B2", result)

    def test_embedded_json_reloaded_coord_map_remaps_addresses(self):
        coord_map = json.loads(json.dumps(ps.build_coord_map([2], [3])))
        int_key = _fmt_key("integer", "0")
        encoding = {
            "cells": {},
            "formats": {int_key: ["C2"]},
            "coord_map": coord_map,
        }
        result = ps.to_paper_compressed_prompt(encoding)
        self.assertEqual(result, "(IntNum|A1)")

    def test_non_compressible_format_key_still_emits_literal(self):
        text_key = _fmt_key("text")
        encoding = {
            "cells": {"hello": ["C3"]},
            "formats": {text_key: ["C3"]},
        }
        result = ps.to_paper_compressed_prompt(encoding)
        self.assertIn("(hello|C3)", result)

    def test_separator_inserted_between_tuples(self):
        encoding = {
            "cells": {"A": ["A1"], "B": ["B1"]},
            "formats": {},
        }
        result = ps.to_paper_compressed_prompt(encoding, separator=",")
        self.assertIn(",", result)

    def test_empty_encoding_returns_empty_string(self):
        result = ps.to_paper_compressed_prompt({"cells": {}, "formats": {}})
        self.assertEqual(result, "")


class TestStage2UncompressedPrompt(unittest.TestCase):

    def setUp(self):
        self._tmpdir = tempfile.mkdtemp()
        self._path = os.path.join(self._tmpdir, "test.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "Name"
        ws["B1"] = "Score"
        ws["A2"] = "Alice"
        ws["B2"] = 95
        ws["A3"] = "Bob"
        ws["B3"] = 87
        wb.save(self._path)

    def tearDown(self):
        import shutil
        shutil.rmtree(self._tmpdir, ignore_errors=True)

    def test_reads_sub_range_pair_string(self):
        result = ps.to_stage2_uncompressed_prompt(self._path, "Data", "A1:B2")
        self.assertEqual(result, "A1,Name|B1,Score|A2,Alice|B2,95")

    def test_reads_single_row(self):
        result = ps.to_stage2_uncompressed_prompt(self._path, "Data", "A1:B1")
        self.assertEqual(result, "A1,Name|B1,Score")

    def test_missing_sheet_raises_key_error(self):
        with self.assertRaises(KeyError):
            ps.to_stage2_uncompressed_prompt(self._path, "NoSuchSheet", "A1:B1")


if __name__ == "__main__":
    unittest.main()
