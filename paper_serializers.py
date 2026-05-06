"""Paper-aligned prompt serializers for SpreadsheetLLM (arXiv:2407.09025).

Three canonical formats:

1. ``to_paper_vanilla_prompt`` — ``A1,Year|A2,Profit`` row-major pairs (the
   paper's vanilla detection prompt and Stage 2 uncompressed format).
2. ``to_paper_compressed_prompt`` — ``(Year|A1)``, ``(IntNum|A1:B3)`` tuples
   with semantic labels substituted for compressible numeric/date/email
   regions, and addresses optionally remapped to a compact coordinate space.
3. ``to_stage2_uncompressed_prompt`` — pair-string for an arbitrary sub-range
   of an original workbook, used by Chain-of-Spreadsheet Stage 2.
"""
from __future__ import annotations

import json
import logging
import re
from typing import Any, Dict, Iterable, List, Optional, Tuple

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

logger = logging.getLogger(__name__)

# --- Paper vocabulary --------------------------------------------------------

# Maps internal semantic types (from temp_helpers.detect_semantic_type) to the
# labels used in the paper's prompts (Section 3.3.3 + Appendix L).
SEMANTIC_LABEL: Dict[str, str] = {
    "integer": "IntNum",
    "float": "FloatNum",
    "numeric": "FloatNum",
    "date": "DateData",
    "datetime": "DateData",
    "time": "TimeData",
    "year": "YearData",
    "email": "EmailData",
    "scientific_notation": "ScientificNum",
    "percentage": "PercentageNum",
    "currency": "CurrencyNum",
}

COMPRESSIBLE_TYPES = set(SEMANTIC_LABEL.keys())

# Maximum cell count we'll iterate over for an LLM-supplied or workbook-derived
# range. Pathological inputs like ``A1:ZZZ999999`` would otherwise pin a CPU.
MAX_CELLS_PER_RANGE = 1_000_000


# --- Address helpers ---------------------------------------------------------

_REF_RE = re.compile(r"^([A-Za-z]+)(\d+)$")


def split_ref(ref: str) -> Tuple[int, int]:
    """Split ``A12`` into ``(row=12, col=1)``."""
    match = _REF_RE.match(ref.strip())
    if not match:
        raise ValueError(f"Invalid cell reference: {ref!r}")
    col_letter, row_str = match.group(1), match.group(2)
    return int(row_str), column_index_from_string(col_letter.upper())


def parse_range(rng: str) -> Tuple[int, int, int, int]:
    """Parse ``A1:B3`` (or single ``A1``) into ``(r1, c1, r2, c2)``.

    Reversed endpoints (``C3:A1``) are normalised but logged so callers can
    surface upstream bugs.
    """
    if ":" in rng:
        a, b = rng.split(":", 1)
    else:
        a = b = rng
    r1, c1 = split_ref(a)
    r2, c2 = split_ref(b)
    if r2 < r1 or c2 < c1:
        logger.warning("parse_range: normalising reversed endpoints in %r", rng)
        r1, r2 = sorted((r1, r2))
        c1, c2 = sorted((c1, c2))
    return r1, c1, r2, c2


def _check_range_size(r1: int, c1: int, r2: int, c2: int, label: str) -> None:
    """Raise ``ValueError`` if a range covers more than ``MAX_CELLS_PER_RANGE``
    cells. Use before any nested per-cell iteration over an attacker- or
    LLM-controlled range."""
    cells = (r2 - r1 + 1) * (c2 - c1 + 1)
    if cells > MAX_CELLS_PER_RANGE:
        raise ValueError(
            f"{label}: range covers {cells} cells, exceeds "
            f"MAX_CELLS_PER_RANGE={MAX_CELLS_PER_RANGE}"
        )


def format_ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def format_range(r1: int, c1: int, r2: int, c2: int) -> str:
    if r1 == r2 and c1 == c2:
        return format_ref(r1, c1)
    return f"{format_ref(r1, c1)}:{format_ref(r2, c2)}"


# --- Coordinate remapping ----------------------------------------------------

CoordMap = Dict[str, Dict[int, int]]


def build_coord_map(kept_rows: Iterable[int], kept_cols: Iterable[int]) -> CoordMap:
    """Build original→compact and inverse maps for retained rows/cols.

    The paper performs coordinate remapping after structural-anchor extraction
    so the LLM sees a continuous, compact grid. This helper returns both
    directions so predicted compact ranges can be unmapped back to the
    original workbook.
    """
    rows = sorted(set(kept_rows))
    cols = sorted(set(kept_cols))
    row_map = {orig: i + 1 for i, orig in enumerate(rows)}
    col_map = {orig: i + 1 for i, orig in enumerate(cols)}
    return {
        "rows": row_map,
        "cols": col_map,
        "rows_inv": {v: k for k, v in row_map.items()},
        "cols_inv": {v: k for k, v in col_map.items()},
    }


def normalize_coord_map(coord_map: Optional[Dict[str, Any]]) -> CoordMap:
    """Return a coord_map with integer keys/values after JSON reload.

    JSON object keys are always strings, so a saved encoding turns
    ``{"rows": {1: 1}}`` into ``{"rows": {"1": 1}}``. Remapping requires
    integer lookup keys, so normalize both directions before use.
    """
    normalized: CoordMap = {"rows": {}, "cols": {}, "rows_inv": {}, "cols_inv": {}}
    if not isinstance(coord_map, dict):
        return normalized
    for axis in normalized:
        mapping = coord_map.get(axis, {})
        if not isinstance(mapping, dict):
            continue
        for key, value in mapping.items():
            try:
                normalized[axis][int(key)] = int(value)
            except (TypeError, ValueError):
                logger.warning(
                    "Skipping non-integer coord_map entry %s[%r]=%r",
                    axis,
                    key,
                    value,
                )
    return normalized


def remap_ref(ref: str, coord_map: CoordMap) -> Optional[str]:
    """Apply original→compact map to a single cell. Returns ``None`` if any axis is unmapped."""
    coord_map = normalize_coord_map(coord_map)
    row, col = split_ref(ref)
    new_row = coord_map.get("rows", {}).get(row)
    new_col = coord_map.get("cols", {}).get(col)
    if new_row is None or new_col is None:
        return None
    return format_ref(new_row, new_col)


def remap_range(rng: str, coord_map: CoordMap) -> Optional[str]:
    """Apply original→compact map to a range. Returns ``None`` if any endpoint is unmapped."""
    coord_map = normalize_coord_map(coord_map)
    r1, c1, r2, c2 = parse_range(rng)
    rows = coord_map.get("rows", {})
    cols = coord_map.get("cols", {})
    new_r1, new_c1 = rows.get(r1), cols.get(c1)
    new_r2, new_c2 = rows.get(r2), cols.get(c2)
    if None in (new_r1, new_c1, new_r2, new_c2):
        return None
    return format_range(new_r1, new_c1, new_r2, new_c2)


def unremap_range(rng: str, coord_map: CoordMap) -> Optional[str]:
    """Reverse a coord_map: compact range → original range."""
    coord_map = normalize_coord_map(coord_map)
    inverted = {
        "rows": coord_map.get("rows_inv", {}),
        "cols": coord_map.get("cols_inv", {}),
    }
    return remap_range(rng, inverted)


# --- Vanilla prompt ----------------------------------------------------------

def _escape_value(val: Any) -> str:
    """Convert a cell value to text safe for the paper's pipe-delimited format."""
    if val is None:
        return ""
    return str(val).replace("|", " ").replace("\n", " ").replace("\r", " ")


def to_paper_vanilla_prompt(sheet) -> str:
    """Encode an openpyxl worksheet as ``A1,value|A2,value|...`` row-major.

    Empty cells are emitted as ``A1,`` (matching the paper's "A1, |A2,Profit"
    examples). All cells in the worksheet bounding box are included so the
    layout remains structurally intact.
    """
    parts: List[str] = []
    max_row = sheet.max_row or 0
    max_col = sheet.max_column or 0
    if max_row and max_col:
        _check_range_size(1, 1, max_row, max_col, "to_paper_vanilla_prompt")
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            ref = format_ref(r, c)
            val = sheet.cell(row=r, column=c).value
            parts.append(f"{ref},{_escape_value(val)}")
    return "|".join(parts)


def to_vanilla_prompt_for_workbook(workbook_path: str) -> Dict[str, str]:
    """Return ``{sheet_name: vanilla_prompt}`` for every sheet in ``workbook_path``."""
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    return {name: to_paper_vanilla_prompt(wb[name]) for name in wb.sheetnames}


# --- Compressed prompt -------------------------------------------------------

def label_for_format_key(format_key: str) -> Optional[str]:
    """Return the paper label for a serialised ``{type, nfs}`` key, or ``None``
    if the type is not compressible (text/boolean/etc.)."""
    try:
        info = json.loads(format_key)
    except (TypeError, ValueError):
        return None
    sem_type = info.get("type")
    if sem_type in COMPRESSIBLE_TYPES:
        return SEMANTIC_LABEL[sem_type]
    return None


def _row_major_sort_key(token: str) -> Tuple[int, int]:
    match = re.search(r"\|([A-Z]+)(\d+)", token)
    if not match:
        return (0, 0)
    return (int(match.group(2)), column_index_from_string(match.group(1)))


def to_paper_compressed_prompt(
    sheet_encoding: Dict[str, Any],
    coord_map: Optional[CoordMap] = None,
    separator: str = "",
) -> str:
    """Render a single-sheet encoding as paper-faithful compressed tuples.

    Steps:

    1. Identify cells covered by an aggregated *compressible* format region
       (numeric, date, email, etc.) — those are emitted once as
       ``(Label|range)``.
    2. Remaining literal-value cells in ``sheet_encoding["cells"]`` are
       emitted as ``(value|address)`` or ``(value|range)`` for merged ranges.
    3. Coordinates are optionally remapped to a compact continuous space via
       ``coord_map``.

    ``separator`` is inserted between tuples; default is empty (paper-style
    ``(Year|A1)( |B1)…`` concatenation).
    """
    coord_map_input = coord_map or sheet_encoding.get("coord_map")
    coord_map = normalize_coord_map(coord_map_input) if coord_map_input else None

    cells: Dict[str, List[str]] = sheet_encoding.get("cells", {}) or {}
    formats: Dict[str, List[str]] = sheet_encoding.get("formats", {}) or {}

    # Step 1: collect label regions and the cell coordinates they cover.
    label_regions: List[Tuple[str, str]] = []
    covered: set = set()
    for fmt_key, ranges in formats.items():
        label = label_for_format_key(fmt_key)
        if label is None:
            continue
        for rng in ranges or []:
            try:
                r1, c1, r2, c2 = parse_range(rng)
            except ValueError:
                continue
            label_regions.append((label, rng))
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    covered.add((r, c))

    # Step 2: literal value tuples from `cells`, minus covered cells.
    literal_tuples: List[Tuple[str, str]] = []  # (range_or_ref, value)
    for value, ranges in cells.items():
        for rng in ranges or []:
            try:
                r1, c1, r2, c2 = parse_range(rng)
            except ValueError:
                continue
            range_cells = [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
            uncovered = [pt for pt in range_cells if pt not in covered]
            if not uncovered:
                continue
            if len(uncovered) == len(range_cells):
                literal_tuples.append((rng, value))
            else:
                for r, c in uncovered:
                    literal_tuples.append((format_ref(r, c), value))

    rendered: List[str] = []
    for label, rng in label_regions:
        out_range = remap_range(rng, coord_map) if coord_map else rng
        if out_range is None:
            continue
        rendered.append(f"({label}|{out_range})")
    for rng, value in literal_tuples:
        out_range = remap_range(rng, coord_map) if coord_map else rng
        if out_range is None:
            continue
        rendered.append(f"({_escape_value(value)}|{out_range})")

    rendered.sort(key=_row_major_sort_key)
    return separator.join(rendered)


# --- Stage 2 uncompressed for sub-range -------------------------------------

def to_stage2_uncompressed_prompt(
    workbook_path: str,
    sheet_name: str,
    table_range: str,
) -> str:
    """Read ``workbook_path`` and produce a paper-vanilla pair-string for
    just ``table_range`` of ``sheet_name``.

    Used by Chain-of-Spreadsheet Stage 2, which the paper requires to be
    uncompressed (Section 4.2). The range must be in original-workbook
    coordinates — callers holding compact ranges should ``unremap_range``
    first.
    """
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet '{sheet_name}' not in workbook {workbook_path}")
    sheet = wb[sheet_name]
    r1, c1, r2, c2 = parse_range(table_range)
    _check_range_size(r1, c1, r2, c2, "to_stage2_uncompressed_prompt")
    parts: List[str] = []
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ref = format_ref(r, c)
            val = sheet.cell(row=r, column=c).value
            parts.append(f"{ref},{_escape_value(val)}")
    return "|".join(parts)


def stage2_pairs_from_rows(
    workbook_path: str,
    sheet_name: str,
    table_range: str,
    rows: Iterable[int],
) -> str:
    """Stage 2 pair-string for an arbitrary subset of rows within a column band.

    Used by ``table_split_qa`` to emit ``header_rows + chunk_rows`` per call
    while keeping the column extent of the identified table.
    """
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Sheet '{sheet_name}' not in workbook {workbook_path}")
    sheet = wb[sheet_name]
    _, c1, _, c2 = parse_range(table_range)
    parts: List[str] = []
    for r in rows:
        for c in range(c1, c2 + 1):
            ref = format_ref(r, c)
            val = sheet.cell(row=r, column=c).value
            parts.append(f"{ref},{_escape_value(val)}")
    return "|".join(parts)


__all__ = [
    "SEMANTIC_LABEL",
    "COMPRESSIBLE_TYPES",
    "split_ref",
    "parse_range",
    "format_ref",
    "format_range",
    "build_coord_map",
    "normalize_coord_map",
    "remap_ref",
    "remap_range",
    "unremap_range",
    "to_paper_vanilla_prompt",
    "to_vanilla_prompt_for_workbook",
    "label_for_format_key",
    "to_paper_compressed_prompt",
    "to_stage2_uncompressed_prompt",
    "stage2_pairs_from_rows",
]
