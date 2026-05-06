"""End-to-end integration tests for coord_map within spreadsheet_llm_encode."""
import os
import sys
import tempfile
import shutil
import unittest

import openpyxl
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(__file__))
from Spreadsheet_LLM_Encoder import spreadsheet_llm_encode
import paper_serializers as ps


def _make_workbook(path: str) -> str:
    """Create an xlsx that produces structural anchors (bold headers + numeric data block).

    Bold font on row 1 and the gap at row 5 cause find_structural_anchors to
    identify boundaries, so kept_rows/cols are non-empty and the encoding has
    actual cells + formats.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Name";  ws["A1"].font = Font(bold=True)
    ws["B1"] = "Score"; ws["B1"].font = Font(bold=True)
    ws["A2"] = "Alice"; ws["B2"] = 95;  ws["B2"].number_format = "0"
    ws["A3"] = "Bob";   ws["B3"] = 72;  ws["B3"].number_format = "0"
    ws["A4"] = "Carol"; ws["B4"] = 88;  ws["B4"].number_format = "0"
    # Row 5 is blank — creates a boundary so row 6 becomes a new anchor
    ws["A6"] = "Total"; ws["A6"].font = Font(bold=True)
    ws["B6"] = 255
    wb.save(path)
    return path


class TestCoordMapIntegration(unittest.TestCase):

    def setUp(self):
        self._tmpdir = tempfile.mkdtemp()
        self._path = _make_workbook(os.path.join(self._tmpdir, "test.xlsx"))
        self._result = spreadsheet_llm_encode(self._path)

    def tearDown(self):
        shutil.rmtree(self._tmpdir, ignore_errors=True)

    def test_encode_returns_sheets_key(self):
        self.assertIn("sheets", self._result)

    def test_sheet1_present_in_encoding(self):
        self.assertIn("Sheet1", self._result["sheets"])

    def test_coord_map_key_exists(self):
        sheet_enc = self._result["sheets"]["Sheet1"]
        self.assertIn("coord_map", sheet_enc)

    def test_coord_map_has_required_sub_keys(self):
        coord_map = self._result["sheets"]["Sheet1"]["coord_map"]
        for key in ("rows", "cols", "rows_inv", "cols_inv"):
            self.assertIn(key, coord_map,
                          msg=f"coord_map missing expected key '{key}'")

    def test_coord_map_rows_and_cols_are_dicts(self):
        coord_map = self._result["sheets"]["Sheet1"]["coord_map"]
        self.assertIsInstance(coord_map["rows"], dict)
        self.assertIsInstance(coord_map["cols"], dict)
        self.assertIsInstance(coord_map["rows_inv"], dict)
        self.assertIsInstance(coord_map["cols_inv"], dict)

    def test_coord_map_rows_non_empty(self):
        coord_map = self._result["sheets"]["Sheet1"]["coord_map"]
        self.assertGreater(len(coord_map["rows"]), 0)

    def test_coord_map_inverse_consistency(self):
        """For every mapping orig->compact, the inverse should map compact->orig."""
        coord_map = self._result["sheets"]["Sheet1"]["coord_map"]
        for orig, compact in coord_map["rows"].items():
            self.assertEqual(coord_map["rows_inv"][compact], orig)
        for orig, compact in coord_map["cols"].items():
            self.assertEqual(coord_map["cols_inv"][compact], orig)

    def test_remap_range_a1_returns_string_or_none(self):
        """A1 may or may not be in the retained rows/cols; either outcome is valid."""
        coord_map = self._result["sheets"]["Sheet1"]["coord_map"]
        result = ps.remap_range("A1", coord_map)
        self.assertTrue(result is None or isinstance(result, str),
                        msg=f"Expected str or None, got {result!r}")

    def test_retained_row1_maps_to_valid_compact_ref(self):
        """Row 1, col 1 is structural (header) so likely retained."""
        coord_map = self._result["sheets"]["Sheet1"]["coord_map"]
        if 1 in coord_map["rows"] and 1 in coord_map["cols"]:
            compact = ps.remap_range("A1", coord_map)
            self.assertIsNotNone(compact)
            # Must be a valid ref parseable back to a positive row/col
            r, c = ps.split_ref(compact)
            self.assertGreater(r, 0)
            self.assertGreater(c, 0)
        else:
            self.skipTest("A1 not in retained cells for this workbook")

    def test_compressed_prompt_with_explicit_coord_map_is_non_empty(self):
        """Passing coord_map explicitly should produce a non-empty tuple string."""
        sheet_enc = self._result["sheets"]["Sheet1"]
        coord_map = sheet_enc["coord_map"]
        prompt = ps.to_paper_compressed_prompt(sheet_enc, coord_map=coord_map)
        self.assertIsInstance(prompt, str)
        self.assertGreater(len(prompt), 0)

    def test_compressed_prompt_auto_coord_map_is_non_empty(self):
        """When coord_map is embedded in sheet_enc, auto-pickup should apply remapping."""
        sheet_enc = self._result["sheets"]["Sheet1"]
        # to_paper_compressed_prompt picks up coord_map from sheet_enc automatically
        prompt = ps.to_paper_compressed_prompt(sheet_enc)
        self.assertIsInstance(prompt, str)
        self.assertGreater(len(prompt), 0)

    def test_unremap_round_trips_a_retained_compact_ref(self):
        """Pick the first retained row/col, remap then unremap — must restore original."""
        coord_map = self._result["sheets"]["Sheet1"]["coord_map"]
        if not coord_map["rows"] or not coord_map["cols"]:
            self.skipTest("No retained rows/cols")
        orig_row = next(iter(coord_map["rows"]))
        orig_col = next(iter(coord_map["cols"]))
        original_ref = ps.format_ref(orig_row, orig_col)
        compact = ps.remap_range(original_ref, coord_map)
        self.assertIsNotNone(compact)
        restored = ps.unremap_range(compact, coord_map)
        self.assertEqual(restored, original_ref)


if __name__ == "__main__":
    unittest.main()
