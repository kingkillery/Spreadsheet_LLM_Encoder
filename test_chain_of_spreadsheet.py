import os
import tempfile
import unittest
import warnings
from unittest.mock import patch

import openpyxl

import chain_of_spreadsheet as cos
import paper_serializers
from llm_backend import EchoBackend


class TestChainOfSpreadsheet(unittest.TestCase):

    def setUp(self):
        # Reset module-level state so tests are isolated.
        cos._BACKEND = None
        cos._STAGE2_LEGACY_WARNED = False

        self.sample_encoding = {
            "sheets": {
                "Sheet1": {
                    "cells": {
                        "Header1": ["A1"],
                        "Value1": ["A2"],
                    }
                }
            }
        }
        self.sample_query = "What is the value?"

    # ------------------------------------------------------------------
    # Stage 1: identify_table
    # ------------------------------------------------------------------

    @patch('chain_of_spreadsheet._call_llm')
    def test_identify_table(self, mock_call_llm):
        mock_call_llm.return_value = "['range': 'A1:B2']"

        table_range = cos.identify_table(self.sample_encoding, self.sample_query)

        # LLM must have been called exactly once.
        mock_call_llm.assert_called_once()
        prompt_arg = mock_call_llm.call_args[0][0]

        # Query is in the prompt.
        self.assertIn(self.sample_query, prompt_arg)

        # Stage-1 now uses paper-compressed tuples, NOT raw JSON.
        # The setUp encoding has no coord_map so addresses stay original.
        self.assertIn("(Header1|A1)", prompt_arg)

        # Response was parsed into the range correctly.
        self.assertEqual(table_range, "A1:B2")

    # ------------------------------------------------------------------
    # Stage 2: generate_response
    # ------------------------------------------------------------------

    @patch('chain_of_spreadsheet._call_llm')
    def test_generate_response(self, mock_call_llm):
        mock_call_llm.return_value = "[C5]"
        sheet_data = self.sample_encoding["sheets"]["Sheet1"]

        response = cos.generate_response(sheet_data, self.sample_query)

        mock_call_llm.assert_called_once()
        prompt_arg = mock_call_llm.call_args[0][0]
        self.assertIn(self.sample_query, prompt_arg)
        self.assertEqual(response, "[C5]")

    # ------------------------------------------------------------------
    # table_split_qa — legacy path (no workbook_path)
    # ------------------------------------------------------------------

    @patch('chain_of_spreadsheet.generate_response')
    def test_table_split_qa_small_table(self, mock_generate_response):
        mock_generate_response.return_value = "Direct answer"

        response = cos.table_split_qa(
            self.sample_encoding["sheets"]["Sheet1"],
            "A1:A2",
            self.sample_query,
        )

        # No workbook_path → legacy path → single generate_response call.
        mock_generate_response.assert_called_once()
        self.assertEqual(response, "Direct answer")

    def test_table_split_qa_large_table_no_workbook_warns_and_calls_once(self):
        """Without workbook_path, table_split_qa warns and calls generate_response once."""
        large_encoding = {"cells": {"a" * 5000: ["A1"]}}

        with patch('chain_of_spreadsheet.generate_response') as mock_gr:
            mock_gr.return_value = "Sub-answer"
            with self.assertLogs('chain_of_spreadsheet', level='WARNING') as cm:
                response = cos.table_split_qa(
                    large_encoding, "A1:Z100", self.sample_query
                )

        # Called exactly once — no chunking without a workbook.
        mock_gr.assert_called_once()
        self.assertEqual(response, "Sub-answer")
        # A deprecation warning must have been emitted in the logs.
        self.assertTrue(
            any("deprecated" in line.lower() or "legacy" in line.lower() for line in cm.output),
            msg=f"Expected a deprecation/legacy warning; got: {cm.output}",
        )

    # ------------------------------------------------------------------
    # table_split_qa — real chunking path
    # ------------------------------------------------------------------

    def test_table_split_qa_real_chunking(self):
        """With workbook_path+sheet_name, real row-chunking is applied."""
        # Build a temp workbook: 30 rows × 4 cols of small numeric data.
        tmp_dir = tempfile.mkdtemp()
        tmp_xlsx = os.path.join(tmp_dir, "test_chunk.xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        for r in range(1, 31):
            for c in range(1, 5):
                ws.cell(row=r, column=c, value=r * 10 + c)
        wb.save(tmp_xlsx)

        # Identity-style coord_map: rows 1-30, cols 1-4 map to themselves.
        coord_map = paper_serializers.build_coord_map(range(1, 31), range(1, 5))

        # EchoBackend records every call.
        backend = EchoBackend(response="X")
        cos.configure_backend(backend)

        try:
            result = cos.table_split_qa(
                sheet_data={"coord_map": coord_map},
                table_range="A1:D30",
                query="Q",
                workbook_path=tmp_xlsx,
                sheet_name="Sheet",
                coord_map=coord_map,
                token_limit=200,
            )
        finally:
            cos.configure_backend(None)
            os.remove(tmp_xlsx)
            os.rmdir(tmp_dir)

        # Real chunking: chunk calls plus a final synthesis call.
        self.assertGreater(len(backend.calls), 1)
        self.assertIn("Candidate Answers", backend.calls[-1])
        self.assertEqual(result, "X")

    # ------------------------------------------------------------------
    # table_split_qa — over-budget single row still dispatched with a warning
    # ------------------------------------------------------------------

    def test_table_split_qa_single_row_exceeds_budget(self):
        """When a single header+row prompt exceeds token_limit, the chunker
        emits a warning and still dispatches the chunk (rather than hanging
        or silently dropping rows)."""
        tmp_dir = tempfile.mkdtemp()
        tmp_xlsx = os.path.join(tmp_dir, "tiny.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        # Two rows, two cols. Even header+1 row will exceed token_limit=1.
        ws['A1'] = "Header"
        ws['B1'] = "Header"
        ws['A2'] = "Value"
        ws['B2'] = "Value"
        wb.save(tmp_xlsx)

        coord_map = paper_serializers.build_coord_map(range(1, 3), range(1, 3))
        backend = EchoBackend(response="ans")
        cos.configure_backend(backend)
        try:
            with self.assertLogs('chain_of_spreadsheet', level='WARNING') as cm:
                result = cos.table_split_qa(
                    sheet_data={"coord_map": coord_map},
                    table_range="A1:B2",
                    query="Q",
                    workbook_path=tmp_xlsx,
                    sheet_name="Sheet",
                    coord_map=coord_map,
                    token_limit=1,
                )
        finally:
            cos.configure_backend(None)
            os.remove(tmp_xlsx)
            os.rmdir(tmp_dir)

        self.assertGreaterEqual(len(backend.calls), 2)
        self.assertIn("Candidate Answers", backend.calls[-1])
        self.assertEqual(result, "ans")
        self.assertTrue(
            any("token_limit" in line for line in cm.output),
            msg=f"Expected over-budget warning; got: {cm.output}",
        )

    # ------------------------------------------------------------------
    # _call_llm error guard
    # ------------------------------------------------------------------

    def test_call_llm_raises_not_implemented(self):
        """_call_llm must raise NotImplementedError until a backend is configured."""
        with self.assertRaises(NotImplementedError):
            cos._call_llm("any prompt")

    def test_range_regex_accepts_double_quotes(self):
        """LLMs sometimes emit double-quoted ranges; the parser must accept both."""
        match = cos._RANGE_RE.search('["range": "A1:F9"]')
        self.assertIsNotNone(match)
        self.assertEqual(match.group(1), "A1:F9")

        match = cos._RANGE_RE.search("['range': 'B2:Z10']")
        self.assertIsNotNone(match)
        self.assertEqual(match.group(1), "B2:Z10")


if __name__ == '__main__':
    unittest.main()
