"""
synthesize_qa.py
----------------
For each annotated XLSX in an annotated_dir (produced by synthesize_benchmark.py),
generate QA pairs and write a companion JSON file compatible with
evaluation.py:load_qa_dataset.

Usage:
    python scripts/synthesize_qa.py <annotated_dir> [--per-table 4] [--seed 42]
         [--out-suffix _qa]
"""

import argparse
import json
import random
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _cell_ref(row, col_idx):
    """Return Excel cell ref, e.g. 'B5'."""
    return f"{get_column_letter(col_idx)}{row}"


def _range_to_parts(range_str):
    """Parse 'A1:F9' -> (r1, c1, r2, c2) as ints."""
    left, right = range_str.split(":")
    col_start = "".join(c for c in left if c.isalpha())
    row_start = "".join(c for c in left if c.isdigit())
    col_end = "".join(c for c in right if c.isalpha())
    row_end = "".join(c for c in right if c.isdigit())
    return (
        int(row_start),
        column_index_from_string(col_start),
        int(row_end),
        column_index_from_string(col_end),
    )


def _is_numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


# ---------------------------------------------------------------------------
# Question generators
# ---------------------------------------------------------------------------


def q_cell_lookup(rng, ws, r1, c1, r2, c2):
    """Pick a random non-header body cell and ask for its value."""
    # Header row is r1; body starts r1+1
    if r2 <= r1:
        return None
    row = rng.randint(r1 + 1, r2)
    col = rng.randint(c1, c2)

    # Get column header name (from row r1)
    header_cell = ws.cell(row=r1, column=col)
    col_name = header_cell.value if header_cell.value is not None else get_column_letter(col)

    # Get row label from first column
    row_label_cell = ws.cell(row=row, column=c1)
    row_label = row_label_cell.value if row_label_cell.value is not None else str(row)

    ref = _cell_ref(row, col)
    question = f"What is the value in the '{col_name}' column for '{row_label}'?"
    answer = f"[{ref}]"
    return {"question": question, "answer": answer}


def q_column_aggregation(rng, ws, r1, c1, r2, c2):
    """Pick a numeric column and ask sum/avg/min/max."""
    if r2 <= r1:
        return None

    # Find numeric columns
    numeric_cols = []
    for col in range(c1, c2 + 1):
        for row in range(r1 + 1, r2 + 1):
            val = ws.cell(row=row, column=col).value
            if _is_numeric(val):
                numeric_cols.append(col)
                break

    if not numeric_cols:
        return None

    col = rng.choice(numeric_cols)
    agg = rng.choice(["SUM", "AVG", "MIN", "MAX"])

    header_cell = ws.cell(row=r1, column=col)
    col_name = header_cell.value if header_cell.value is not None else get_column_letter(col)

    col_letter = get_column_letter(col)
    data_start = r1 + 1
    data_end = r2
    ref = f"{col_letter}{data_start}:{col_letter}{data_end}"

    question = f"What is the {agg} of the '{col_name}' column from row {data_start} to row {data_end}?"
    answer = f"[{agg}({ref})]"
    return {"question": question, "answer": answer}


def q_cross_column_comparison(rng, ws, r1, c1, r2, c2):
    """Ask which row has the largest value in a numeric column."""
    if r2 <= r1:
        return None

    numeric_cols = []
    for col in range(c1, c2 + 1):
        for row in range(r1 + 1, r2 + 1):
            val = ws.cell(row=row, column=col).value
            if _is_numeric(val):
                numeric_cols.append(col)
                break

    if not numeric_cols:
        return None

    col = rng.choice(numeric_cols)
    header_cell = ws.cell(row=r1, column=col)
    col_name = header_cell.value if header_cell.value is not None else get_column_letter(col)

    # Find the row with the max value in this column
    best_row = r1 + 1
    best_val = None
    for row in range(r1 + 1, r2 + 1):
        val = ws.cell(row=row, column=col).value
        if _is_numeric(val):
            if best_val is None or val > best_val:
                best_val = val
                best_row = row

    ref = _cell_ref(best_row, col)
    question = f"Which row has the largest value in the '{col_name}' column?"
    answer = f"[{ref}]"
    return {"question": question, "answer": answer}


def q_header_lookup(rng, ws, r1, c1, r2, c2):
    """Ask what the column name is for a given column letter."""
    col = rng.randint(c1, c2)
    col_letter = get_column_letter(col)
    ref = _cell_ref(r1, col)
    question = f"What is the column name for column {col_letter}?"
    answer = f"[{ref}]"
    return {"question": question, "answer": answer}


QUESTION_TYPES = [
    q_cell_lookup,
    q_column_aggregation,
    q_cross_column_comparison,
    q_header_lookup,
]

# ---------------------------------------------------------------------------
# Per-file QA generation
# ---------------------------------------------------------------------------


def generate_qa_for_file(xlsx_path, annotation, per_table, rng):
    """Return a list of QA pair dicts for one XLSX file."""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    all_pairs = []

    tables = annotation.get("tables", [])
    for table_meta in tables:
        sheet_name = table_meta.get("sheet", "Sheet1")
        range_str = table_meta.get("range", "")
        if not range_str:
            continue

        try:
            r1, c1, r2, c2 = _range_to_parts(range_str)
        except Exception:
            continue

        if sheet_name not in wb.sheetnames:
            ws = wb.active
        else:
            ws = wb[sheet_name]

        # Cycle through question types
        generated = 0
        attempt = 0
        max_attempts = per_table * len(QUESTION_TYPES) * 2

        while generated < per_table and attempt < max_attempts:
            q_fn = QUESTION_TYPES[attempt % len(QUESTION_TYPES)]
            pair = q_fn(rng, ws, r1, c1, r2, c2)
            attempt += 1
            if pair is not None:
                all_pairs.append(pair)
                generated += 1

    return all_pairs


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main():
    parser = argparse.ArgumentParser(
        description="Generate QA pairs for annotated XLSX files."
    )
    parser.add_argument("annotated_dir", help="Directory containing annotated XLSX+JSON pairs")
    parser.add_argument("--per-table", type=int, default=4,
                        help="Number of QA pairs to generate per table")
    parser.add_argument("--seed", type=int, default=42, help="Random seed")
    parser.add_argument("--out-suffix", default="_qa",
                        help="Suffix appended to base filename for QA JSON output")
    args = parser.parse_args()

    annotated_dir = Path(args.annotated_dir)
    rng = random.Random(args.seed)

    written = 0
    for xlsx_path in sorted(annotated_dir.glob("*.xlsx")):
        # Load matching annotation JSON (same stem, no suffix)
        ann_path = xlsx_path.with_suffix(".json")
        if not ann_path.exists():
            continue

        with open(ann_path, "r") as f:
            annotation = json.load(f)

        qa_pairs = generate_qa_for_file(xlsx_path, annotation, args.per_table, rng)
        if not qa_pairs:
            continue

        out_json = xlsx_path.parent / (xlsx_path.stem + args.out_suffix + ".json")
        output = {
            "qa_pairs": qa_pairs,
            "tables": annotation.get("tables", []),
        }
        with open(out_json, "w") as f:
            json.dump(output, f, indent=2)

        written += 1

    print(f"Wrote {written} files to {annotated_dir}")


if __name__ == "__main__":
    main()
